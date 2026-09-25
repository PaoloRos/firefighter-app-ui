"""Verification for the version-controlled XLSX sample schedule."""

from datetime import date, datetime, time
from importlib.metadata import version
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from firefighter_tools_backend.services.calendar_conversion import (
    convert_calendar,
)

ASSET_PATH = (
    Path(__file__).parents[2]
    / "assets"
    / "examples"
    / "calendar_schedule_example.xlsx"
)
EXPECTED_COLUMNS = [
    "id",
    "summary",
    "all_date",
    "start_date",
    "start_time",
    "end_date",
    "end_time",
    "location",
    "description",
    "participants",
]
EXPECTED_IDS = [
    "dienst-2026-08-08",
    "uebung-2026-08-12",
    "nachtdienst-2026-08-15",
]


def test_sample_is_small_and_stored_in_application_assets() -> None:
    assert ASSET_PATH.is_file()
    assert ASSET_PATH.relative_to(Path(__file__).parents[2]).parts[:2] == (
        "assets",
        "examples",
    )
    assert ASSET_PATH.stat().st_size < 100 * 1024


def test_sample_uses_exact_v030_columns_and_native_cell_types() -> None:
    workbook = load_workbook(ASSET_PATH, data_only=True)
    try:
        assert workbook.sheetnames == ["Beispiel"]
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    assert list(rows[0]) == EXPECTED_COLUMNS
    assert len(rows) == len(EXPECTED_IDS) + 1
    assert all(len(row) == len(EXPECTED_COLUMNS) for row in rows)
    assert [row[0] for row in rows[1:]] == EXPECTED_IDS
    assert all(isinstance(row[3], (date, datetime)) for row in rows[1:])
    assert all(isinstance(row[5], (date, datetime)) for row in rows[1:])

    # The all-day entry leaves both time cells genuinely empty rather than a
    # zero time; the timed entries carry native time cells.
    assert [row[2] for row in rows[1:]] == [True, False, False]
    assert [row[4] for row in rows[1:]] == [None, time(19, 0), time(22, 0)]
    assert [row[6] for row in rows[1:]] == [None, time(21, 0), time(7, 0)]

    # One event for everyone (empty cell), one shared, one personal.
    assert [row[9] for row in rows[1:]] == [None, "101;204", "204"]


def test_sample_personal_conversion_keeps_only_that_persons_events() -> None:
    with ASSET_PATH.open("rb") as source:
        result = convert_calendar(
            source,
            filename=ASSET_PATH.name,
            participant="101",
        )

    assert result.total_count == 2
    assert "UID:dienst-2026-08-08" in result.ics_text
    assert "UID:uebung-2026-08-12" in result.ics_text
    assert "UID:nachtdienst-2026-08-15" not in result.ics_text


def test_sample_converts_all_events_with_calendar_conversion_v030() -> None:
    assert version("calendar-conversion") == "0.3.0"

    with ASSET_PATH.open("rb") as source:
        result = convert_calendar(source, filename=ASSET_PATH.name)

    assert (
        result.total_count,
        result.converted_count,
        result.skipped_count,
    ) == (len(EXPECTED_IDS), len(EXPECTED_IDS), 0)
    assert result.invalid_events == ()
    for event_id in EXPECTED_IDS:
        assert f"UID:{event_id}" in result.ics_text
    assert "SUMMARY:Bereitschaftsdienst" in result.ics_text
    assert "LOCATION:Feuerwehrhaus" in result.ics_text
    # An all-day entry must emit DATE values, never a midnight timestamp.
    assert "DTSTART;VALUE=DATE:20260808" in result.ics_text
    assert "DTEND;VALUE=DATE:20260809" in result.ics_text


def test_sample_is_downloaded_from_one_stable_application_url(
    client: TestClient,
) -> None:
    response = client.get(
        "/api/v1/tools/calendar-converter/example",
    )

    assert response.status_code == 200
    assert response.content == ASSET_PATH.read_bytes()
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="calendar_schedule_example.xlsx"'
    )


def test_endpoint_accepts_the_sample_as_a_successful_upload(
    client: TestClient,
) -> None:
    with ASSET_PATH.open("rb") as source:
        response = client.post(
            "/api/v1/tools/calendar-converter/convert",
            files={
                "file": (
                    ASSET_PATH.name,
                    source,
                    (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )
            },
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "success"
    assert (
        payload["total_count"],
        payload["converted_count"],
        payload["skipped_count"],
    ) == (len(EXPECTED_IDS), len(EXPECTED_IDS), 0)
    assert payload["invalid_events"] == []
    assert payload["calendar"]["filename"] == "calendar_schedule_example.ics"
